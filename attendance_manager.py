from openpyxl import load_workbook
from datetime import datetime
import os


EXCEL_FILE = "attendance/attendance.xlsx"


def mark_attendance(name, roll, department):

    # Excel file না থাকলে error দেখাবে
    if not os.path.exists(EXCEL_FILE):
        print("Attendance Excel file not found!")
        return "Excel file not found"

    # Excel open
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook["Attendance"]

    # Current date and time
    current_date = datetime.now().strftime("%Y-%m-%d")
    current_time = datetime.now().strftime("%H:%M:%S")

    # Check today's attendance
    for row in sheet.iter_rows(min_row=2, values_only=True):

        date = row[0]
        existing_roll = row[3]

        if date == current_date and str(existing_roll) == str(roll):

            workbook.close()

            print(f"{name} is already marked Present today.")

            return "Already Present"

    # New attendance entry
    sheet.append([
        current_date,
        current_time,
        name,
        roll,
        department,
        "Present"
    ])

    workbook.save(EXCEL_FILE)
    workbook.close()

    print("--------------------------------")
    print("Attendance Marked Successfully!")
    print("Name       :", name)
    print("Roll       :", roll)
    print("Department :", department)
    print("Date       :", current_date)
    print("Time       :", current_time)
    print("--------------------------------")

    return "Attendance Marked"