from flask import Flask, render_template, request, jsonify

import os
import json
import base64
import cv2
import numpy as np

from deepface import DeepFace

app = Flask(__name__)


# ==============================
# Dashboard
# ==============================

@app.route("/")
def home():

    from openpyxl import load_workbook
    from datetime import datetime, date

    # -----------------------------
    # Total Students + Registered Rolls
    # -----------------------------

    students_folder = "students"

    registered_rolls = set()

    if os.path.exists(students_folder):

        for filename in os.listdir(students_folder):

            if filename.endswith(".json"):

                file_path = os.path.join(students_folder, filename)

                try:
                    with open(file_path, "r") as file:
                        student = json.load(file)

                    roll = str(student.get("roll", "")).strip()

                    if roll:
                        registered_rolls.add(roll)

                except Exception as e:
                    print("Error reading student:", filename)
                    print(e)

    total_students = len(registered_rolls)

    # -----------------------------
    # Today's Attendance
    # -----------------------------

    today = date.today()

    today_present_rolls = set()

    attendance_file = "attendance/attendance.xlsx"

    recent_records = []

    if os.path.exists(attendance_file):

        workbook = load_workbook(attendance_file)

        sheet = workbook.active

        for row in sheet.iter_rows(min_row=2, values_only=True):

            if row[0] is None:
                continue

            excel_date = row[0]
            roll = str(row[3]).strip() if row[3] is not None else ""
            status = str(row[5]).strip().lower() if row[5] is not None else ""

            # -----------------------------
            # Check Today's Date
            # -----------------------------

            is_today = False

            if isinstance(excel_date, datetime):

                is_today = excel_date.date() == today

            elif isinstance(excel_date, date):

                is_today = excel_date == today

            else:

                date_value = str(excel_date).strip()

                if date_value.startswith(today.strftime("%Y-%m-%d")):
                    is_today = True

            # -----------------------------
            # Count Unique Registered Students
            # -----------------------------

            if (
                is_today
                and status == "present"
                and roll in registered_rolls
            ):
                today_present_rolls.add(roll)

            # -----------------------------
            # Recent Records
            # -----------------------------

            recent_records.append({
                "date": row[0],
                "time": row[1],
                "name": row[2],
                "roll": row[3],
                "department": row[4],
                "status": row[5]
            })

    # -----------------------------
    # Total Present
    # -----------------------------

    today_present = len(today_present_rolls)

    # Safety check
    if today_present > total_students:
        today_present = total_students

    # -----------------------------
    # Today's Absent
    # -----------------------------

    today_absent = total_students - today_present

    # -----------------------------
    # Attendance Rate
    # -----------------------------

    if total_students > 0:

        attendance_rate = round(
            (today_present / total_students) * 100,
            2
        )

    else:

        attendance_rate = 0

    # -----------------------------
    # Recent 5 Records
    # -----------------------------

    recent_records = recent_records[-5:]

    recent_records.reverse()

    # -----------------------------
    # Dashboard
    # -----------------------------

    return render_template(
        "index.html",
        total_students=total_students,
        today_present=today_present,
        today_absent=today_absent,
        attendance_rate=attendance_rate,
        recent_records=recent_records
    )

    # Last 5 attendance records
    recent_records = recent_records[-5:]

    # -----------------------------
    # Absent
    # -----------------------------

    today_absent = total_students - today_present

    if today_absent < 0:
        today_absent = 0

    # -----------------------------
    # Attendance Rate
    # -----------------------------

    if total_students > 0:

        attendance_rate = round(
            (today_present / total_students) * 100,
            2
        )

    else:

        attendance_rate = 0

    return render_template(
        "index.html",
        total_students=total_students,
        today_present=today_present,
        today_absent=today_absent,
        attendance_rate=attendance_rate,
        recent_records=recent_records
    )

# ==============================
# Register Student
# ==============================

@app.route("/register")
def register():
    return render_template("register.html")


# ==============================
# Take Attendance
# ==============================

@app.route("/attendance")
def attendance():
    return render_template("attendance.html")


# ==============================
# Registered Students
# ==============================

@app.route("/students")
def students():

    students_list = []

    students_folder = "students"

    if os.path.exists(students_folder):

        for filename in os.listdir(students_folder):

            if filename.endswith(".json"):

                file_path = os.path.join(
                    students_folder,
                    filename
                )

                try:

                    with open(file_path, "r") as file:
                        student = json.load(file)

                    students_list.append(student)

                except Exception as e:

                    print("Error reading:", filename)
                    print(e)

    return render_template(
        "students.html",
        students=students_list
    )


# ==============================
# Attendance Records
# ==============================

#new add 10.41/////////////////////////////////////////////////////////////////////////

@app.route("/records")
def records():

    from openpyxl import load_workbook

    attendance_records = []

    attendance_file = "attendance/attendance.xlsx"

    # Search/filter values
    search_name = request.args.get("name", "").strip().lower()
    search_roll = request.args.get("roll", "").strip()
    search_date = request.args.get("date", "").strip()
    search_department = request.args.get("department", "").strip().lower()

    if os.path.exists(attendance_file):

        workbook = load_workbook(attendance_file)
        sheet = workbook.active

        for row in sheet.iter_rows(min_row=2, values_only=True):

            if row[0] is None:
                continue

            date_value = str(row[0])
            time_value = str(row[1])
            name_value = str(row[2])
            roll_value = str(row[3])
            department_value = str(row[4])
            status_value = str(row[5])

            # -----------------------------
            # Apply Filters
            # -----------------------------

            if search_name:
                if search_name not in name_value.lower():
                    continue

            if search_roll:
                if search_roll not in roll_value:
                    continue

            if search_date:
                if search_date != date_value:
                    continue

            if search_department:
                if search_department not in department_value.lower():
                    continue

            attendance_records.append({
                "date": date_value,
                "time": time_value,
                "name": name_value,
                "roll": roll_value,
                "department": department_value,
                "status": status_value
            })

    return render_template(
        "records.html",
        records=attendance_records,
        search_name=search_name,
        search_roll=search_roll,
        search_date=search_date,
        search_department=search_department
    )



# ==============================
# Register Student from Web
# ==============================

@app.route("/register_student", methods=["POST"])
def register_student():

    try:

        data = request.get_json()

        name = data.get("name")
        roll = data.get("roll")
        department = data.get("department")
        image_data = data.get("image")

        if not name or not roll or not department or not image_data:

            return jsonify({
                "success": False,
                "message": "All information is required."
            })

        os.makedirs("students", exist_ok=True)

        # Convert Base64 image to OpenCV image

        image_data = image_data.split(",")[1]

        image_bytes = base64.b64decode(image_data)

        image_array = np.frombuffer(
            image_bytes,
            np.uint8
        )

        frame = cv2.imdecode(
            image_array,
            cv2.IMREAD_COLOR
        )

        if frame is None:

            return jsonify({
                "success": False,
                "message": "Could not process captured image."
            })

        # Create Face Embedding

        embedding_result = DeepFace.represent(
            img_path=frame,
            model_name="Facenet512",
            detector_backend="opencv",
            enforce_detection=True
        )

        if len(embedding_result) == 0:

            return jsonify({
                "success": False,
                "message": "No face detected."
            })

        face_embedding = embedding_result[0]["embedding"]

        # Student Data

        student_data = {
            "name": name,
            "roll": roll,
            "department": department,
            "face_embedding": face_embedding
        }

        # Save Student JSON

        file_path = f"students/{roll}.json"

        with open(file_path, "w") as file:

            json.dump(
                student_data,
                file,
                indent=4
            )

        print("\n================================")
        print("Student registered successfully!")
        print("================================")
        print("Name       :", name)
        print("Roll       :", roll)
        print("Department :", department)
        print("Saved file :", file_path)

        return jsonify({

            "success": True,

            "message": (
                f"Student registered successfully! "
                f"Name: {name}, Roll: {roll}"
            )

        })

    except Exception as e:

        print("\nRegistration error:")
        print(e)

        return jsonify({

            "success": False,

            "message": f"Registration failed: {str(e)}"

        })


# ==============================
# Recognize Face from Web Camera
# ==============================

@app.route("/recognize_face", methods=["POST"])
def recognize_face():

    try:

        # Get image from browser

        data = request.get_json()

        image_data = data.get("image")

        if not image_data:

            return jsonify({
                "success": False,
                "message": "No image received."
            })

        # Convert Base64 image to OpenCV image

        image_data = image_data.split(",")[1]

        image_bytes = base64.b64decode(image_data)

        image_array = np.frombuffer(
            image_bytes,
            np.uint8
        )

        frame = cv2.imdecode(
            image_array,
            cv2.IMREAD_COLOR
        )

        if frame is None:

            return jsonify({
                "success": False,
                "message": "Could not process image."
            })

        # Create current face embedding

        embedding_result = DeepFace.represent(
            img_path=frame,
            model_name="Facenet512",
            detector_backend="opencv",
            enforce_detection=True
        )

        if len(embedding_result) == 0:

            return jsonify({
                "success": False,
                "message": "No face detected."
            })

        current_embedding = np.array(
            embedding_result[0]["embedding"]
        )

        # Search registered students

        students_folder = "students"

        if not os.path.exists(students_folder):

            return jsonify({
                "success": False,
                "message": "No registered students found."
            })

        recognized_student = None

        best_distance = float("inf")

        # Compare with saved embeddings

        for filename in os.listdir(students_folder):

            if not filename.endswith(".json"):
                continue

            file_path = os.path.join(
                students_folder,
                filename
            )

            try:

                with open(file_path, "r") as file:

                    student = json.load(file)

                saved_embedding = np.array(
                    student["face_embedding"]
                )

                # Cosine similarity

                cosine_similarity = np.dot(
                    current_embedding,
                    saved_embedding
                ) / (
                    np.linalg.norm(current_embedding)
                    *
                    np.linalg.norm(saved_embedding)
                )

                distance = 1 - cosine_similarity

                if distance < best_distance:

                    best_distance = distance

                    recognized_student = student

            except Exception as e:

                print(
                    "Error reading:",
                    filename,
                    e
                )

        # Recognition threshold

        threshold = 0.30

        if (
            recognized_student is None
            or best_distance > threshold
        ):

            return jsonify({

                "success": False,

                "message": "Unknown face."

            })

        # Mark attendance

        from attendance_manager import mark_attendance

        attendance_result = mark_attendance(

            recognized_student["name"],

            recognized_student["roll"],

            recognized_student["department"]

        )

        # Return result

        return jsonify({

            "success": True,

            "message":
                f"{recognized_student['name']} "
                f"({recognized_student['roll']}) - "
                f"{attendance_result}"

        })

    except Exception as e:

        print("\nFace recognition error:")
        print(e)

        return jsonify({

            "success": False,

            "message":
                f"Recognition failed: {str(e)}"

        })

#new add korci 8.22........................................................................


@app.route("/edit_student/<roll>", methods=["GET", "POST"])
def edit_student(roll):

    file_path = f"students/{roll}.json"

    if not os.path.exists(file_path):
        return "Student not found."

    with open(file_path, "r") as file:
        student = json.load(file)

    if request.method == "POST":

        name = request.form.get("name")
        department = request.form.get("department")

        if not name or not department:
            return "Name and Department are required."

        student["name"] = name
        student["department"] = department

        with open(file_path, "w") as file:
            json.dump(student, file, indent=4)

        return """
        <script>
            alert("Student information updated successfully!");
            window.location.href = "/students";
        </script>
        """

    return render_template(
        "edit_student.html",
        student=student
    )


@app.route("/delete_student/<roll>", methods=["POST"])
def delete_student(roll):

    file_path = f"students/{roll}.json"

    if not os.path.exists(file_path):
        return jsonify({
            "success": False,
            "message": "Student not found."
        })

    try:

        os.remove(file_path)

        return jsonify({
            "success": True,
            "message": "Student deleted successfully."
        })

    except Exception as e:

        return jsonify({
            "success": False,
            "message": str(e)
        })

#new add korci 8.22---------------------------------------------------------------------------------


# ==============================
# Start Flask
# ==============================

if __name__ == "__main__":

    app.run(debug=True)






